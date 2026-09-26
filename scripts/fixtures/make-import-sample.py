"""Generate src/modules/wiki/lib/__fixtures__/import-sample.pptx for the PowerPoint import tests.

Usage (python-pptx 1.0.x in a throwaway virtualenv):
    python3 -m venv /tmp/pptx-venv && /tmp/pptx-venv/bin/pip install python-pptx
    /tmp/pptx-venv/bin/python scripts/fixtures/make-import-sample.py

The deck is 16:9. The default python-pptx template is 4:3, so the master and layout
placeholders are stretched horizontally by 4/3 to match; slide placeholders keep no xfrm
of their own and inherit geometry from layout/master, which is what the importer must
resolve. The expected geometry asserted in presentation-pptx.test.ts is printed at the end.
"""
import io
import os
import struct
import zlib

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

OUT = os.path.join(os.path.dirname(__file__), "..", "..", "src", "modules", "wiki", "lib", "__fixtures__", "import-sample.pptx")


def png(width, height, rgb):
    """A tiny solid-colour PNG without needing Pillow."""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    chunk = lambda kind, data: struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
stretch = 12192000 / 9144000
for owner in [prs.slide_master, *prs.slide_layouts]:
    for placeholder in owner.placeholders:
        if not placeholder._element.xpath("./p:spPr/a:xfrm"):
            continue  # inherits from the master, which is stretched itself
        placeholder.left, placeholder.width = Emu(round(placeholder.left * stretch)), Emu(round(placeholder.width * stretch))

# 1: title slide, both placeholders inherit their geometry from the layout.
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Quarterly Review"
slide.placeholders[1].text = "Management Platform 2026"

# 2: title and content with bullets, one of them nested.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Roadmap"
body = slide.placeholders[1].text_frame
body.text = "Launch presentation import"
for text, level in [("Improve the canvas editor", 0), ("Frames as containers", 1), ("Ship collaboration", 0)]:
    paragraph = body.add_paragraph()
    paragraph.text, paragraph.level = text, level

# 3: filled shape, image and table.
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Budget Overview"
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(2))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x2E, 0x86, 0xDE)
box.line.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
box.line.width = Pt(1.5)
box.text_frame.text = "Total 120k"
slide.shapes.add_picture(io.BytesIO(png(8, 8, (0xE6, 0x7E, 0x22))), Inches(1), Inches(4.5), Inches(1), Inches(1))
table = slide.shapes.add_table(3, 3, Inches(5), Inches(2), Inches(6), Inches(1.5)).table
for row, values in enumerate([["Item", "Q1", "Q2"], ["Staff", "40", "45"], ["Tools", "10", "12"]]):
    for column, value in enumerate(values):
        table.cell(row, column).text = value

# 4: bar chart with speaker notes.
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Revenue by Quarter"
data = CategoryChartData()
data.categories = ["Q1", "Q2", "Q3", "Q4"]
data.add_series("Revenue", (12, 18, 27, 21))
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(2), Inches(8), Inches(4.5), data)
slide.notes_slide.notes_text_frame.text = "Emphasise the Q3 growth.\nMention the new customers."

prs.save(OUT)

for index, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        print(index, shape.shape_type, shape.name, shape.left, shape.top, shape.width, shape.height)
